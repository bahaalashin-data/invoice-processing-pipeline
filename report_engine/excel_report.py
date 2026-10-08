# Builds the Excel report: All Invoices, Summary Statistics (formula-driven), Rejected Invoices

from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import REPORTS_DIR
from logger_config import get_logger

logger = get_logger(__name__)

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True)
BODY_FONT = Font(name=FONT_NAME)
CURRENCY_FMT = "$#,##0.00"

ALL_INVOICES_HEADERS = [
    "Invoice No",
    "Source Batch",
    "Invoice Date",
    "Created Date",
    "Customer Name",
    "Product Category",
    "Product Name",
    "Quantity",
    "Unit Price",
    "Total Amount",
    "Country",
]
REJECTED_HEADERS = ["Source File", "Invoice No", "Source Batch", "Customer Name", "Product Name", "Reason"]


def _style_header_row(ws, headers: List[str], row: int = 1) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _autosize_columns(ws, widths: List[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_all_invoices_sheet(wb: Workbook, valid_records: List[dict]) -> int:
    ws = wb.active
    ws.title = "All Invoices"
    _style_header_row(ws, ALL_INVOICES_HEADERS)

    for row_idx, inv in enumerate(valid_records, start=2):
        ws.cell(row=row_idx, column=1, value=inv["invoice_no"]).font = BODY_FONT
        ws.cell(row=row_idx, column=2, value=inv.get("source_batch", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=3, value=inv["invoice_date"]).font = BODY_FONT
        ws.cell(row=row_idx, column=4, value=inv["created_date"]).font = BODY_FONT
        ws.cell(row=row_idx, column=5, value=inv["customer_name"]).font = BODY_FONT
        ws.cell(row=row_idx, column=6, value=inv["product_category"]).font = BODY_FONT
        ws.cell(row=row_idx, column=7, value=inv["product_name"]).font = BODY_FONT
        ws.cell(row=row_idx, column=8, value=inv["quantity"]).font = BODY_FONT

        c7 = ws.cell(row=row_idx, column=9, value=inv["unit_price"])
        c7.font = BODY_FONT
        c7.number_format = CURRENCY_FMT

        c8 = ws.cell(row=row_idx, column=10, value=inv["total_amount"])
        c8.font = BODY_FONT
        c8.number_format = CURRENCY_FMT

        ws.cell(row=row_idx, column=11, value=inv["country"]).font = BODY_FONT

    _autosize_columns(ws, [16, 24, 14, 14, 22, 18, 20, 10, 12, 14, 14])
    ws.freeze_panes = "A2"
    return len(valid_records)


def _write_rejected_sheet(wb: Workbook, rejected_records: List[dict]) -> int:
    ws = wb.create_sheet("Rejected Invoices")
    _style_header_row(ws, REJECTED_HEADERS)

    for row_idx, rec in enumerate(rejected_records, start=2):
        ws.cell(row=row_idx, column=1, value=rec.get("source_file", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=2, value=rec.get("invoice_no", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=3, value=rec.get("source_batch", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=4, value=rec.get("customer_name", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=5, value=rec.get("product_name", "")).font = BODY_FONT
        ws.cell(row=row_idx, column=6, value=rec.get("rejection_reason") or rec.get("reason", "")).font = BODY_FONT

    _autosize_columns(ws, [26, 16, 24, 22, 20, 55])
    ws.freeze_panes = "A2"
    return len(rejected_records)



def _write_summary_sheet(wb: Workbook, invoice_row_count: int, rejected_row_count: int) -> None:
    ws = wb.create_sheet("Summary Statistics")

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22

    inv_last_row = invoice_row_count + 1
    rej_last_row = max(rejected_row_count + 1, 2)

    ws["A1"] = "Summary Statistics"
    ws["A1"].font = TITLE_FONT

    # KPI Block

    kpi_rows = [
        (
            "Total Invoices",
            f"=COUNTA('All Invoices'!A2:A{inv_last_row})"
        ),

        (
            "Total Revenue",
            f"=SUM('All Invoices'!J2:J{inv_last_row})"
        ),

        (
            "Total Customers",
            f"=IFERROR(SUMPRODUCT(1/COUNTIF('All Invoices'!E2:E{inv_last_row},"
            f"'All Invoices'!E2:E{inv_last_row})),0)"
        ),

        (
            "Total Quantity Sold",
            f"=SUM('All Invoices'!H2:H{inv_last_row})"
        ),

        (
            "Average Invoice Value",
            f"=IFERROR(AVERAGE('All Invoices'!J2:J{inv_last_row}),0)"
        ),

        (
            "Rejected Records",
            f"=COUNTA('Rejected Invoices'!A2:A{rej_last_row})"
        ),
    ]

    row = 3

    for label, formula in kpi_rows:

        ws.cell(
            row=row,
            column=1,
            value=label
        ).font = LABEL_FONT

        cell = ws.cell(
            row=row,
            column=2,
            value=formula
        )

        cell.font = BODY_FONT

        if "Revenue" in label or "Average" in label:
            cell.number_format = CURRENCY_FMT

        row += 1

    # Revenue By Category

    row += 1

    categories = sorted(
        {
            r
            for r in _column_values(
                wb,
                "All Invoices",
                6,      # Product Category
                inv_last_row,
            )
        }
    )

    ws.cell(
        row=row,
        column=1,
        value="Revenue by Category"
    ).font = LABEL_FONT

    row += 1

    for category in categories:

        ws.cell(
            row=row,
            column=1,
            value=category
        ).font = BODY_FONT

        cell = ws.cell(
            row=row,
            column=2,
            value=(
                f"=SUMIF("
                f"'All Invoices'!F2:F{inv_last_row},"
                f"A{row},"
                f"'All Invoices'!J2:J{inv_last_row}"
                f")"
            ),
        )

        cell.font = BODY_FONT
        cell.number_format = CURRENCY_FMT

        row += 1

    # Revenue By Customer

    row += 1

    customers = sorted(
        {
            r
            for r in _column_values(
                wb,
                "All Invoices",
                5,      # Customer Name
                inv_last_row,
            )
        }
    )

    ws.cell(
        row=row,
        column=1,
        value="Revenue by Customer"
    ).font = LABEL_FONT

    row += 1

    for customer in customers:

        ws.cell(
            row=row,
            column=1,
            value=customer
        ).font = BODY_FONT

        cell = ws.cell(
            row=row,
            column=2,
            value=(
                f"=SUMIF("
                f"'All Invoices'!E2:E{inv_last_row},"
                f"A{row},"
                f"'All Invoices'!J2:J{inv_last_row}"
                f")"
            ),
        )

        cell.font = BODY_FONT
        cell.number_format = CURRENCY_FMT

        row += 1

    ws.sheet_view.showGridLines = False

def _column_values(wb: Workbook, sheet_name: str, col_idx: int, last_row: int):
    ws = wb[sheet_name]
    for row in range(2, last_row + 1):
        value = ws.cell(row=row, column=col_idx).value
        if value not in (None, ""):
            yield value


def generate_excel_report(
    valid_records: List[dict],
    rejected_records: List[dict],
    output_path: Optional[Path] = None,
) -> Path:
    """
    Builds Invoice_Report.xlsx with 3 sheets and returns its path.
    Sheet order: All Invoices -> Summary Statistics -> Rejected Invoices
    (matches the order data flows through the pipeline).
    """
    wb = Workbook()

    invoice_count = _write_all_invoices_sheet(wb, valid_records)
    rejected_count = _write_rejected_sheet(wb, rejected_records)
    _write_summary_sheet(wb, invoice_count, rejected_count)

    # Reorder sheets: All Invoices, Summary Statistics, Rejected Invoices
    wb.move_sheet("Summary Statistics", offset=-1)

    if output_path is None:
        output_path = Path(REPORTS_DIR) / "Invoice_Report.xlsx"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)
    logger.info(f"Excel report generated -> {output_path} ({invoice_count} invoices, {rejected_count} rejected)")
    return output_path
